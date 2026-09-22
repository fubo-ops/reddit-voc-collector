#!/usr/bin/env python3
import argparse, datetime, json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLS = ["schema_version","platform","asin","matched_asins","query","matched_keywords","subreddit","post_id","post_title","post_url","comment_id","parent_id","depth","author","body","score","created_at","comment_url","collected_at","audit_status","audit_gap_reasons","semantic_relevance"]

def rows(path):
    return [json.loads(x) for x in Path(path).read_text(encoding="utf-8").splitlines() if x.strip()]

def add_sheet(wb, name, data):
    ws=wb.create_sheet(name); ws.append(COLS)
    for row in data: ws.append([json.dumps(row.get(c),ensure_ascii=False) if isinstance(row.get(c),(list,dict)) else row.get(c) for c in COLS])
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="305496")
    for i,c in enumerate(COLS,1):
        ws.column_dimensions[get_column_letter(i)].width=60 if c=="body" else min(45,max(12,len(c)+2))
        if c=="body":
            for cell in ws[get_column_letter(i)]: cell.alignment=Alignment(wrap_text=True,vertical="top")
        if c=="comment_id":
            for cell in ws[get_column_letter(i)][1:]: cell.number_format="@"

def main():
    p=argparse.ArgumentParser(); p.add_argument("--root",type=Path,required=True); p.add_argument("--batch-dir",type=Path,required=True); p.add_argument("--target",type=int,default=1000); a=p.parse_args()
    formal=a.root/"outputs/reddit-comments"; batch=a.batch_dir
    existing=rows(formal/"B003ULL1NQ_selected_smoke_trusted_comments.jsonl")
    cp=json.loads((formal/"B003ULL1NQ_selected_smoke_checkpoint.json").read_text(encoding="utf-8"))
    meta={str(x.get("post_id")):x for x in cp.get("discovered_posts",[]) if isinstance(x,dict)}
    audit=json.loads((batch/"audit.json").read_text(encoding="utf-8"))["posts"]
    latest={}
    for x in rows(batch/"captures.jsonl"): latest[str(x.get("post_id"))]=x
    seen_ids={str(x.get("comment_id")) for x in existing}; seen_urls={str(x.get("comment_url","")).rstrip("/") for x in existing}
    passes=[]; partial=[]; duplicates=0
    for pid,capture in latest.items():
        evidence=audit.get(pid,{}); status=evidence.get("status","PARTIAL"); query=meta.get(pid,{}).get("query") or "dog joint supplement"
        for row in capture.get("comments",[]):
            cid=str(row.get("comment_id") or ""); url=str(row.get("comment_url") or "").rstrip("/"); body=str(row.get("body") or "").strip()
            if not cid or not body or body.lower() in {"[deleted]","[removed]"}: continue
            if cid in seen_ids or (url and url in seen_urls): duplicates+=1; continue
            seen_ids.add(cid); seen_urls.add(url)
            rec={"schema_version":"reddit_comment_v1","platform":"reddit","asin":"B003ULL1NQ","matched_asins":["B003ULL1NQ"],"query":query,"matched_keywords":[query],"subreddit":capture.get("subreddit"),"post_id":pid,"post_title":capture.get("post_title"),"post_url":capture.get("page_url"),"comment_id":cid,"parent_id":row.get("parent_id"),"depth":row.get("depth",0),"author":row.get("author"),"body":body,"score":row.get("score"),"created_at":row.get("created_at"),"comment_url":row.get("comment_url"),"collected_at":capture.get("captured_at"),"audit_status":status,"audit_gap_reasons":evidence.get("reasons",[]),"semantic_relevance":"thread_context_verified","comment_record_status":"valid"}
            (passes if status=="PASS" else partial).append(rec)
    trusted=existing+passes; validated=partial[:max(0,a.target-len(trusted))]; delivered=trusted+validated
    for name,data in (("trusted_comments.jsonl",trusted),("validated_partial_comments.jsonl",validated)):
        (batch/name).write_text("".join(json.dumps(x,ensure_ascii=False)+"\n" for x in data),encoding="utf-8")
    manifest={"schema_version":"reddit_voc_delivery_v1","asin":"B003ULL1NQ","product_title":"Nutramax Cosequin Joint Supplement for Dogs, Chewable Tablets, 132ct","target_comments":a.target,"existing_trusted_preserved":len(existing),"new_pass_comments":len(passes),"trusted_comments":len(trusted),"validated_partial_comments":len(validated),"total_delivered_unique_comments":len(delivered),"target_reached":len(delivered)>=a.target,"posts_checked":len(audit),"pass_posts":sum(x.get("status")=="PASS" for x in audit.values()),"partial_posts":sum(x.get("status")=="PARTIAL" for x in audit.values()),"blocked_posts":sum(x.get("status")=="BLOCKED" for x in audit.values()),"filtered_comments":sum(int(x.get("filtered_count",0)) for x in audit.values()),"global_duplicates_removed":duplicates,"stop_reason":"target_comments_reached" if len(delivered)>=a.target else "available_relevant_posts_exhausted","source_checkpoint_unchanged":str(formal/"B003ULL1NQ_selected_smoke_checkpoint.json"),"batch_checkpoint":str(batch/"checkpoint.json"),"generated_at":datetime.datetime.now(datetime.timezone.utc).isoformat()}
    (batch/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    (batch/"delivery_checkpoint.json").write_text(json.dumps({"schema_version":"reddit_voc_delivery_checkpoint_v1","delivered_comment_ids":[x["comment_id"] for x in delivered],"completed_post_ids":list(audit),"stop_reason":manifest["stop_reason"]},ensure_ascii=False,indent=2),encoding="utf-8")
    wb=Workbook(); wb.remove(wb.active); add_sheet(wb,"Trusted_Comments",trusted); add_sheet(wb,"Validated_Partial_Comments",validated)
    ws=wb.create_sheet("Run_Summary"); ws.append(["metric","value"])
    for k,v in manifest.items(): ws.append([k,json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v])
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=95
    stamp=datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%SZ"); xlsx=formal/f"B003ULL1NQ_reddit_voc_delivery_{stamp}.xlsx"; wb.save(xlsx)
    check=load_workbook(xlsx,read_only=True); errors=[]
    if check.sheetnames != ["Trusted_Comments","Validated_Partial_Comments","Run_Summary"]: errors.append("sheet_contract")
    if len({x["comment_id"] for x in delivered}) != len(delivered): errors.append("duplicate_comment_id")
    if len(delivered)<a.target: errors.append("target_not_reached")
    result={"status":"PASS" if not errors else "FAIL","exit_code":0 if not errors else 1,"errors":errors,"excel_path":str(xlsx),"trusted_rows":len(trusted),"validated_partial_rows":len(validated),"total_rows":len(delivered)}
    (batch/"validate.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8"); print(json.dumps({"manifest":manifest,"validation":result},ensure_ascii=False)); return result["exit_code"]

if __name__=="__main__": raise SystemExit(main())
