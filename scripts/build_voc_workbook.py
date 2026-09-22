#!/usr/bin/env python3
"""Build the six-sheet Reddit conversation-discovery VOC workbook."""
from __future__ import annotations
import argparse, json, re
from pathlib import Path

SHEETS=("Trusted_Comments","Partial_Candidates","Conversation_Map","Query_Performance","Community_Map","Run_Summary")
ILLEGAL=re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

def rows(path):
    if not path or not path.exists(): return []
    if path.suffix.lower()==".jsonl": return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]
    value=json.loads(path.read_text(encoding="utf-8")); return value if isinstance(value,list) else [value]
def obj(path,default):
    return json.loads(path.read_text(encoding="utf-8")) if path and path.exists() else default

def flatten_map(cmap):
    out=[]
    for layer,groups in (cmap.get("layers") or {}).items():
        for group,values in groups.items():
            for value in values if isinstance(values,list) else [values]:
                if value not in (None,""): out.append({"layer":layer,"group":group,"value":value,"source_asin":cmap.get("asin")})
    return out

def clean(value):
    if isinstance(value,(list,dict)): value=json.dumps(value,ensure_ascii=False)
    return ILLEGAL.sub("",str(value if value is not None else ""))

def add_sheet(wb,name,data):
    ws=wb.create_sheet(name); data=list(data or [])
    fields=list(dict.fromkeys(k for row in data for k in row)) or ["status"]
    ws.append(fields)
    for row in data: ws.append([clean(row.get(f)) for f in fields])
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    from openpyxl.styles import Alignment, Font
    for cell in ws[1]: cell.font=Font(bold=True)
    for column in ws.columns:
        letter=column[0].column_letter; ws.column_dimensions[letter].width=min(60,max(12,max(len(str(c.value or "")) for c in column)+2))
        for cell in column: cell.alignment=Alignment(vertical="top",wrap_text=True)
    return ws

def build(output,trusted,partial,cmap,queries,communities,summary):
    try: from openpyxl import Workbook
    except ImportError as exc: raise SystemExit("openpyxl is required; use the Codex bundled Python runtime") from exc
    wb=Workbook(); wb.remove(wb.active)
    add_sheet(wb,"Trusted_Comments",trusted); add_sheet(wb,"Partial_Candidates",partial)
    add_sheet(wb,"Conversation_Map",flatten_map(cmap)); add_sheet(wb,"Query_Performance",queries)
    add_sheet(wb,"Community_Map",communities); add_sheet(wb,"Run_Summary",[{"metric":k,"value":v} for k,v in summary.items()])
    output.parent.mkdir(parents=True,exist_ok=True); wb.save(output); return output

def main():
    p=argparse.ArgumentParser(description=__doc__)
    for flag in ("trusted","partial","conversation-map","query-plan","community-map","manifest"): p.add_argument("--"+flag,type=Path)
    p.add_argument("--output",type=Path,required=True); a=p.parse_args()
    cmap=obj(a.conversation_map,{}); plan=obj(a.query_plan,{}); community=obj(a.community_map,[]); summary=obj(a.manifest,{})
    build(a.output,rows(a.trusted),rows(a.partial),cmap,plan.get("queries",[]),community if isinstance(community,list) else [community],summary)
    print(json.dumps({"status":"written","output":str(a.output.resolve()),"sheets":SHEETS},ensure_ascii=False)); return 0
if __name__=="__main__": raise SystemExit(main())
